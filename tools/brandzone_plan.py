# -*- coding: utf-8 -*-
"""브랜드존 배정안 생성기.

사용자 기준
  ① 피부과·약국 등 전문 채널 중심 공급 (일반 뷰티샵 중심 브랜드는 후순위)
  ② 약국 판매 소구점이 한 문장으로 설명되는가
  ③ 브랜드 자체 인지도 (병행 검토)

배정 원칙
  · 엔드캡(EL·ER)만 브랜드 블록으로 쓴다. 교차통로를 향해 정면 노출되고,
    현재 충전율이 4%로 가장 비어 있으며, 본진 진열을 밀어내지 않는다.
  · 모듈 내 S4·S5 확장 여유는 그대로 둔다. 충전율 11% 상태에서는
    선반을 더 채우는 것보다 기존 품목의 페이싱을 늘리는 것이 먼저다.
  · S1(최상단)은 개구 100mm·상품높이 80mm로 상품 진열에 부적합하므로 안내 유지.
"""
import openpyxl, json, collections, datetime

XLSX = 'data/Pharmacy_Full_Zoning_Shelf_Operations_v2.xlsx'
LAYOUT = 'data/uploaded-integrated-layout-v2.json'

# 채널 3 = 제약·의료기기 계열 또는 병의원·약국 전용 라인
# 채널 2 = 더마 포지셔닝이나 일반 유통 병행
# 채널 1 = 일반 뷰티 채널 중심 (사용자 기준상 후순위)
BRAND_PROFILE = {
    "케어리브":      dict(ch=3, pt=3, maker="니치반", claim="습윤 밴드·드레싱"),
    "프로캄":        dict(ch=3, pt=3, maker="한미사이언스", claim="EGF·멜라·아크페어"),
    "닥터리쥬올":     dict(ch=2, pt=3, maker="Dr.Reju-All", claim="PDRN 스킨케어"),
    "고려은단":       dict(ch=3, pt=2, maker="고려은단", claim="약국 비타민"),
    "리쥬비":        dict(ch=3, pt=3, maker="파마리서치", claim="PDRN 의약품·화장품 구분"),
    "리쥬란":        dict(ch=2, pt=3, maker="파마리서치", claim="PDRN 병의원 계열"),
    "마데카파마시아":  dict(ch=3, pt=3, maker="동국제약", claim="센텔라·재생"),
    "이지덤":        dict(ch=3, pt=3, maker="대웅제약", claim="습윤 드레싱"),
    "아크로패스":     dict(ch=3, pt=3, maker="라파스", claim="마이크로니들 패치"),
    "EGFx":        dict(ch=3, pt=3, maker="대웅제약", claim="EGF 다운타임"),
    "신신":          dict(ch=3, pt=3, maker="신신제약", claim="파스·외용"),
    "정원삼":        dict(ch=3, pt=2, maker="정원삼", claim="홍삼 선물"),
    "에버콜라겐":     dict(ch=2, pt=3, maker="뉴트리", claim="먹는 콜라겐"),
    "락토핏":        dict(ch=2, pt=3, maker="종근당건강", claim="유산균"),
    "엠디스픽":       dict(ch=2, pt=2, maker="확인 필요", claim="더마 기초"),
    "레비온":        dict(ch=2, pt=2, maker="확인 필요", claim="EGF·재생"),
    "닥터알파":       dict(ch=2, pt=2, maker="확인 필요", claim="장벽"),
    "닥터오라클":     dict(ch=2, pt=2, maker="닥터오라클", claim="더마 기초"),
    "마미케어":       dict(ch=2, pt=2, maker="확인 필요", claim="톤·주름"),
    "테라브레스":     dict(ch=2, pt=3, maker="테라브레스", claim="구강 관리"),
    "네츄럴메이드":   dict(ch=2, pt=2, maker="오츠카", claim="영양"),
    "톰":           dict(ch=1, pt=2, maker="톰", claim="PPM·NMN"),
    "비플레인":       dict(ch=1, pt=2, maker="비플레인", claim="일반 뷰티 기초"),
    "네오젠":        dict(ch=1, pt=2, maker="네오젠", claim="레티놀·바쿠치올"),
    "믹순":          dict(ch=1, pt=1, maker="믹순", claim="일반 뷰티"),
    "프랭클리":       dict(ch=1, pt=1, maker="프랭클리", claim="일반 뷰티"),
    "이지듀":        dict(ch=3, pt=3, maker="대웅제약", claim="EGF"),
}

# 엔드캡 → 브랜드 배정. None 은 보류(적합 후보 없음).
ENDCAP_PLAN = {
    "G5-ER": ("리쥬비", "PDRN 의약품(리쥬비넥스·W1)과 화장품(리쥬비-에스)의 구분 안내가 이 엔드의 역할과 정확히 일치"),
    "G5-EL": ("프로캄", "11개 모듈에 걸친 다축 배치. 고민 길찾기 엔드와 정합"),
    "G2-ER": ("케어리브", "선정 SKU 18건으로 최다. G2-ER에 이미 본진 배치"),
    "G2-EL": ("이지덤", "케어리브와 소재·부위 비교 구도를 만든다"),
    "G1-EL": ("아크로패스", "마이크로니들 패치 단일 소구. 여드름 상태 선택 엔드와 정합"),
    "G4-ER": ("고려은단", "약국 채널 비타민 대표. G4-ER에 이미 본진 배치"),
    "G4-EL": ("에버콜라겐", "먹는 뷰티 단일 축"),
    "G3-EL": ("마데카파마시아", "센텔라·재생. G3-D-3 선케어와 연결"),
    "G6-ER": ("신신", "파스·외용 단일 축. L3·G6-U-1 본진과 연결"),
    "G3-ER": (None, "두피·휴식 테마에 부합하는 전문 채널 브랜드 후보 없음. 안내·확장 여유 유지"),
    "G1-ER": (None, "엠디스픽이 9건 배치되어 있으나 제조·채널 미확인. 확인 후 재검토"),
    "G6-EL": (None, "립·눈 테마 전용 전문 채널 브랜드 후보 없음"),
}

wb = openpyxl.load_workbook(XLSX, data_only=True)
sh = [r for r in list(wb['02_선반315'].iter_rows(values_only=True))[6:] if r[0]]
pl = [r for r in list(wb['05_상품별배치367'].iter_rows(values_only=True))[6:] if r[0]]
c3 = {r[0]: r for r in [x for x in list(wb['03_상품분류2526'].iter_rows(values_only=True))[6:] if x[0]]}
zo = {r[0]: r for r in [x for x in list(wb['01_전체조닝63'].iter_rows(values_only=True))[6:] if x[0] and x[0] != '합계']}

placed = collections.defaultdict(list)
for r in pl:
    placed[r[1]].append(r)
brand_of = {}
for r in pl:
    c = c3.get(r[4])
    t = c[19] if c else None
    if t and t != '원본명 참고·브랜드 대조':
        brand_of.setdefault(t, []).append(r)

def eff_w(mod):
    return 645 if mod.endswith(('-EL', '-ER')) else 865

rows, summary = [], []
for mod, (brand, why) in ENDCAP_PLAN.items():
    shelves = sorted([r for r in sh if r[1] == mod], key=lambda x: x[3])
    z = zo[mod]
    cur_f = sum(int(r[16] or 0) for r in shelves)
    cap_f = eff_w(mod) // 70 * len([s for s in shelves if s[3] != 1])
    prof = BRAND_PROFILE.get(brand, {})
    have = brand_of.get(brand, [])
    for s in shelves:
        lv = int(s[3])
        if lv == 1:
            act, note = "유지(안내)", "개구 100mm·상품높이 80mm. 상품 진열 부적합"
        elif brand is None:
            act, note = "유지", why
        elif s[10] == '확장 여유':
            act, note = "브랜드존 신설", f"{brand} 블록"
        else:
            act, note = "브랜드존 편입", f"{brand} 블록(기존 {s[10]} 유지)"
        rows.append(dict(module=mod, shelf=s[0], level=lv, status=s[10],
                         theme=z[2], fixture=z[3], brand=brand or "", action=act, note=note))
    summary.append(dict(module=mod, theme=z[2], fixture=z[3], brand=brand or "—",
                        maker=prof.get("maker", ""), claim=prof.get("claim", ""),
                        channel=prof.get("ch", 0), point=prof.get("pt", 0),
                        placedSku=len(have), placedF=sum(int(r[10] or 0) for r in have),
                        currentF=cur_f, capacityF=cap_f, why=why))

out = {
    "schema": "kpharmacy.brandzone/1.0",
    "generatedAt": datetime.date.today().isoformat(),
    "criteria": {
        "channel": "3=제약·의료기기 계열 또는 병의원·약국 전용 / 2=더마 포지셔닝·일반 유통 병행 / 1=일반 뷰티 채널 중심",
        "point": "3=단일 성분·기능 축으로 한 문장 설명 가능 / 2=라인 정체성 뚜렷 / 1=범용 기초",
        "note": "채널·제조사 판단은 일반 지식에 근거한 분류이며 거래 조건으로 확인해야 함. "
                "04_브랜드검토의 취급 결정은 전 브랜드가 '미검토' 상태.",
    },
    "policy": [
        "엔드캡만 브랜드 블록으로 사용한다. 본진 진열을 밀어내지 않는다.",
        "모듈 내 S4·S5 확장 여유 16칸은 유지한다. 충전율 11%에서는 페이싱 확대가 우선.",
        "S1은 개구 100mm로 상품 진열에 부적합하므로 안내를 유지한다.",
        "브랜드 교체 시에도 고민명과 위치 코드는 유지한다(원안 원칙).",
    ],
    "assignments": summary,
    "shelves": rows,
}
json.dump(out, open('data/brandzone-plan.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

print(f"{'엔드캡':<8}{'주테마':<14}{'브랜드':<12}{'제조·공급':<14}{'채널':>4}{'소구':>4}{'배치SKU':>7}{'현재F':>6}{'가능F':>6}")
for s in summary:
    print(f"{s['module']:<8}{s['theme'][:13]:<14}{s['brand'][:11]:<12}{s['maker'][:13]:<14}"
          f"{s['channel']:>4}{s['point']:>4}{s['placedSku']:>7}{s['currentF']:>6}{s['capacityF']:>6}")
n = sum(1 for s in summary if s['brand'] != '—')
print(f"\n브랜드존 배정 {n}개 엔드캡 / 보류 {len(summary)-n}개")
print(f"대상 선반 {sum(1 for r in rows if '브랜드존' in r['action'])}칸 "
      f"(신설 {sum(1 for r in rows if r['action']=='브랜드존 신설')} · 편입 {sum(1 for r in rows if r['action']=='브랜드존 편입')})")

# ── 레이아웃 반영 ────────────────────────────────────────────────────
layout = json.load(open(LAYOUT, encoding='utf-8'))
S = layout["store"]
els = {e.get("code"): e for e in layout["elements"] if e.get("code")}

# ① 브랜드존 표기
by_mod = collections.defaultdict(list)
for r in rows:
    by_mod[r["module"]].append(r)
applied = 0
for mod, rs in by_mod.items():
    e = els.get(mod)
    if not e:
        continue
    brand = next((r["brand"] for r in rs if r["brand"]), "")
    if not brand:
        continue
    prof = BRAND_PROFILE.get(brand, {})
    z = e.setdefault("zoning", {})
    z["brandZone"] = {
        "brand": brand,
        "maker": prof.get("maker", ""),
        "claim": prof.get("claim", ""),
        "channelTier": prof.get("ch"),
        "pointTier": prof.get("pt"),
        "levels": [r["level"] for r in rs if "브랜드존" in r["action"]],
        "rationale": next(r["note"] for r in rs if r["brand"]),
        "status": "제안 · 거래조건 미확인",
    }
    for r in rs:
        if "브랜드존" not in r["action"]:
            continue
        i = r["level"] - 1           # levelDetails[0] = 최상단 = S1
        if i < len(e.get("levelDetails", [])):
            d = e["levelDetails"][i]
            d["brandZone"] = brand
            d["note"] = f"[브랜드존] {brand} · {prof.get('claim','')} | " + str(d.get("note", ""))
            applied += 1
    e["note"] = f"[브랜드존 제안] {brand} ({prof.get('maker','')}) · {prof.get('claim','')}\n" + str(e.get("note", ""))

# ② 하단 유리문 4,000mm × 2 (사용자 확인 사항)
DOOR_W = 4000
for cid, name, x in (("DR-L", "좌측 유리문 4m", 0), ("DR-R", "우측 유리문 4m", S["w"] - DOOR_W)):
    layout["elements"].append({
        "id": f"door_{cid}", "type": "door", "name": name, "code": cid,
        "x": x, "y": S["h"] - 200, "w": DOOR_W, "h": 200, "zHeight": 2400,
        "bays": 1, "levels": 1, "zone": "공용동선", "category": "출입구",
        "categoryZone": "공용동선", "opacity": 1, "rotation": 0,
        "locked": False, "hidden": False, "levelDetails": [],
        "note": "하단 좌·우 각 4,000mm 개방 유리문(사용자 확인). "
                "현재 B열(x 1,502~9,022)과 폭이 겹쳐 B1~B3·B6~B8 위치 확인 필요.",
    })
# 진입 존을 문 폭에 맞춘다
for e in layout["elements"]:
    if e.get("name") == "좌측 개방 진입":
        e["x"], e["w"] = 0, DOOR_W
    if e.get("name") == "우측 개방 진입":
        e["x"], e["w"] = S["w"] - DOOR_W, DOOR_W

doc = layout["project"].setdefault("zoningDocumentation", {})
doc["revision"] = {
    "date": datetime.date.today().isoformat(),
    "resolved": [
        "X5 기둥 없음(사용자 확인) — 중앙 1,000mm는 통로로 확정. 좌·우 구역 횡단 가능.",
        "하단 좌·우 각 4,000mm 개방 유리문(사용자 확인) — door 요소로 반영.",
    ],
    "added": [
        f"엔드캡 9개에 브랜드존 제안 반영({applied}칸 표기).",
    ],
    "openIssue": [
        "유리문 4,000mm와 B열(1,502~9,022)의 폭이 겹친다. "
        "B열이 고정 유리 뒤 윈도 진열이면 현행 유지, 통로 개방이면 B1~B3·B6~B8 재배치 필요.",
    ],
}
json.dump(layout, open('data/kpharmacy-brandzone-layout.json', 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
print(f"\n레이아웃 반영: 브랜드존 {applied}칸 표기 · 유리문 2개 추가 → data/kpharmacy-brandzone-layout.json")

# ③ 브랜드존 CSV
import csv
with open('data/brandzone-plan.csv', 'w', encoding='utf-8-sig', newline='') as f:
    wcsv = csv.writer(f)
    wcsv.writerow(["엔드캡", "주테마", "매대명", "선반ID", "단", "현재 운영상태",
                   "브랜드", "제조·공급", "소구점", "채널등급", "조치", "비고"])
    for r in rows:
        prof = BRAND_PROFILE.get(r["brand"], {})
        wcsv.writerow([r["module"], r["theme"], r["fixture"], r["shelf"], f"S{r['level']}",
                       r["status"], r["brand"], prof.get("maker", ""), prof.get("claim", ""),
                       prof.get("ch", ""), r["action"], r["note"]])
print("CSV: data/brandzone-plan.csv")
